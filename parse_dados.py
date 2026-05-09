import openpyxl
import json
import os
import unicodedata

def normalize(s):
    """Remove accents and lowercase for fuzzy name matching."""
    return ''.join(c for c in unicodedata.normalize('NFD', str(s).lower()) if unicodedata.category(c) != 'Mn').strip()

wb = openpyxl.load_workbook('data/diagnostico_territorial.xlsx')
ws = wb['DADOS']

records = []
for row in ws.iter_rows(min_row=12, values_only=True):
    if not any(v is not None for v in row):
        continue
    territorio, municipio = row[1], row[2]
    if not municipio:
        continue
    d = {
        'territorio':             str(territorio).strip() if territorio else '',
        'municipio':              str(municipio).strip()  if municipio  else '',
        'area':                   float(row[3])  if row[3]  is not None else None,
        'pop_total':              int(row[4])    if row[4]  is not None else None,
        'pop_urbana':             int(row[5])    if row[5]  is not None else None,
        'pop_rural':              int(row[6])    if row[6]  is not None else None,
        'pop_quilombola':         int(row[7])    if row[7]  is not None else None,
        'terr_quilombolas':       int(row[8])    if row[8]  is not None else None,
        'pop_indigena':           int(row[9])    if row[9]  is not None else None,
        'assentamentos_reforma':  int(row[10])   if row[10] is not None else None,
        'taxa_alfabetizacao':     float(row[11]) if row[11] is not None else None,
        'sexo_hom_pct':           float(row[12]) if row[12] is not None else None,
        'sexo_mul_pct':           float(row[13]) if row[13] is not None else None,
        'idh':                    float(row[14]) if row[14] is not None else None,
        'taxa_crescimento_pop':   float(row[15]) if row[15] is not None else None,
        'pct_extremamente_pobres':float(row[16]) if row[16] is not None else None,
        'benef_agri_familiar':    int(row[17])   if row[17] is not None else None,
        'pessoal_agricultura':    int(row[18])   if row[18] is not None else None,
        'estab_agri_familiar':    int(row[19])   if row[19] is not None else None,
    }
    records.append(d)

print(f'Registros: {len(records)}')
print('Territorios:', sorted(set(r['territorio'] for r in records)))
print()
print('Amostra:', json.dumps(records[0], ensure_ascii=False, indent=2))

with open('data/municipios.json', 'w', encoding='utf-8') as f:
    json.dump(records, f, ensure_ascii=False, indent=2)
print('\nmunicipos.json salvo!')

# === CROSS-REFERENCE IBGE CODES from old xlsx ===
wb_old = openpyxl.load_workbook('data/lista_composicao_territorios_pas_nordeste.xlsx')
ws_old = wb_old['Lista de Municípios']
ibge_map = {}   # normalized_name -> ibge_code (7-digit string)
for row in ws_old.iter_rows(min_row=6, values_only=True):
    ibge_code, muni_name = row[3], row[4]
    if ibge_code and muni_name:
        ibge_map[normalize(muni_name)] = str(int(ibge_code))

# Add ibge to records and report missing
missing = []
for r in records:
    key = normalize(r['municipio'])
    if key in ibge_map:
        r['ibge'] = ibge_map[key]
    else:
        r['ibge'] = None
        missing.append(r['municipio'])

print(f'IBGE matched: {sum(1 for r in records if r["ibge"])} / {len(records)}')
if missing:
    print('Sem match IBGE:', missing)

# Resave with ibge
with open('data/municipios.json', 'w', encoding='utf-8') as f:
    json.dump(records, f, ensure_ascii=False, indent=2)
print('municipios.json atualizado com ibge!')

# === FILTER GEOJSON to the 23 municipalities ===
our_ibge = set(r['ibge'] for r in records if r['ibge'])
with open('data/pb_municipios.geojson', 'r', encoding='utf-8') as f:
    full_geo = json.load(f)

filtered_feats = [f for f in full_geo['features'] if f['properties']['codarea'] in our_ibge]
print(f'GeoJSON features matched: {len(filtered_feats)} / {len(our_ibge)} expected')

new_geo = {'type': 'FeatureCollection', 'features': filtered_feats}
with open('data/pb_municipios.geojson', 'w', encoding='utf-8') as f:
    json.dump(new_geo, f, ensure_ascii=False, separators=(',', ':'))
print('pb_municipios.geojson atualizado!')

# === CAPA info ===
ws_capa = wb['CAPA']
capa = {}
team = []
for row in ws_capa.iter_rows(values_only=True):
    vals = [v for v in row if v is not None]
    if not vals:
        continue
    if vals[0] == 'Equipe:':
        continue
    if len(vals) == 2 and vals[1] not in ('', None):
        team.append({'nome': vals[0], 'papel': vals[1]})
    elif len(vals) == 1:
        capa['fonte'] = vals[0]
capa['equipe'] = team
print('\nCAPA info:', json.dumps(capa, ensure_ascii=False, indent=2))

